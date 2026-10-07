import profile.profile_builder_router as profile_builder_router
import time
from io import BytesIO
from profile.profile_builder import CertificateEntry, Education, SkillEntry, UserProfile, WorkExperience

from docx import Document

from shared.llm_structured import StructuredProviderError


def test_profile_processing_accepts_markdown(client, monkeypatch):
    token = client.post(
        "/api/v1/auth/register", json={"email": "profile-markdown@example.com", "password": "password123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    session = client.post("/api/v1/profile/sessions", headers=headers).json()
    received = {}

    async def capture_profile(document, output_language="English"):
        received.update(document)
        received["output_language"] = output_language
        return UserProfile(Name="Markdown Candidate")

    monkeypatch.setattr(profile_builder_router, "extract_profile", capture_profile)
    upload = client.post(
        f"/api/v1/profile/sessions/{session['id']}/files",
        headers=headers,
        files={"file": ("resume.md", b"# Markdown Candidate\nPython engineer", "text/markdown")},
    )

    assert upload.status_code == 202
    assert received["content"] == "# Markdown Candidate\nPython engineer"
    assert received["name"] == "resume.md"
    assert received["output_language"] == "English"


def test_profile_processing_extracts_docx_text(client, monkeypatch):
    token = client.post(
        "/api/v1/auth/register", json={"email": "profile-docx@example.com", "password": "password123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    session = client.post("/api/v1/profile/sessions", headers=headers).json()
    document = Document()
    document.add_paragraph("DOCX Candidate")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Python engineer"
    content = BytesIO()
    document.save(content)
    received = {}

    async def capture_profile(document, output_language="English"):
        received.update(document)
        received["output_language"] = output_language
        return UserProfile(Name="DOCX Candidate")

    monkeypatch.setattr(profile_builder_router, "extract_profile", capture_profile)
    upload = client.post(
        f"/api/v1/profile/sessions/{session['id']}/files",
        headers=headers,
        files={
            "file": (
                "resume.docx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert upload.status_code == 202
    assert "DOCX Candidate" in received["content"]
    assert "Python engineer" in received["content"]
    assert received["name"] == "resume.docx"
    assert received["output_language"] == "English"


def test_profile_session_language_applies_to_all_extraction(client, monkeypatch):
    token = client.post(
        "/api/v1/auth/register", json={"email": "profile-language@example.com", "password": "password123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    session_response = client.post(
        "/api/v1/profile/sessions",
        headers=headers,
        json={"output_language": "French"},
    )
    session = session_response.json()
    received_languages = []

    async def capture_profile(document, output_language="English"):
        received_languages.append(output_language)
        return UserProfile(Name="Candidat")

    monkeypatch.setattr(profile_builder_router, "extract_profile", capture_profile)
    for filename in ("resume.txt", "experience.txt"):
        upload = client.post(
            f"/api/v1/profile/sessions/{session['id']}/files",
            headers=headers,
            files={"file": (filename, b"Python engineer", "text/plain")},
        )
        assert upload.status_code == 202

    profile = client.get("/api/v1/profile", headers=headers).json()
    session_detail = client.get(f"/api/v1/profile/sessions/{session['id']}", headers=headers).json()
    assert session_response.json()["output_language"] == "French"
    assert session_detail["output_language"] == "French"
    assert received_languages
    assert set(received_languages) == {"French"}
    assert profile["output_language"] == "French"


def test_profile_upload_rejects_unsupported_content_type(client):
    token = client.post(
        "/api/v1/auth/register", json={"email": "profile-unsupported@example.com", "password": "password123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    session = client.post("/api/v1/profile/sessions", headers=headers).json()

    response = client.post(
        f"/api/v1/profile/sessions/{session['id']}/files",
        headers=headers,
        files={"file": ("resume.bin", b"not supported", "application/octet-stream")},
    )

    assert response.status_code == 422


def test_profile_upload_rejects_invalid_docx(client):
    token = client.post(
        "/api/v1/auth/register", json={"email": "profile-invalid-docx@example.com", "password": "password123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    session = client.post("/api/v1/profile/sessions", headers=headers).json()

    response = client.post(
        f"/api/v1/profile/sessions/{session['id']}/files",
        headers=headers,
        files={
            "file": (
                "resume.docx",
                b"not a valid DOCX document",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 422
    assert response.status_code == 422


def test_profile_processing_completes_and_preserves_edit(client):
    token = client.post(
        "/api/v1/auth/register", json={"email": "profile-int@example.com", "password": "password123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    session = client.post("/api/v1/profile/sessions", headers=headers).json()

    upload = client.post(
        f"/api/v1/profile/sessions/{session['id']}/files",
        headers=headers,
        files={"file": ("resume.txt", b"Python engineer", "text/plain")},
    )
    assert upload.status_code == 202
    for _ in range(20):
        profile = client.get("/api/v1/profile", headers=headers)
        if profile.status_code == 200:
            break
        time.sleep(0.01)
    assert profile.status_code == 200

    updated = client.put("/api/v1/profile", headers=headers, json={"data": {"summary": "user confirmed"}})
    assert updated.status_code == 200
    assert updated.json()["data"]["summary"] == "user confirmed"


def test_profile_processing_persists_complete_structured_profile(client, monkeypatch):
    token = client.post(
        "/api/v1/auth/register", json={"email": "profile-complete@example.com", "password": "password123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    session = client.post("/api/v1/profile/sessions", headers=headers).json()

    async def complete_profile(*args, **kwargs):
        return UserProfile(
            Name="Ada Lovelace",
            Email="ada@example.com",
            Phone="+1 555 0100",
            Location="London",
            LinkedIn="https://linkedin.example/ada",
            Summary="Python engineer",
            SoftSkills=["Communication"],
            CoreCompetences=["Technical leadership"],
            AdditionalInformations=["Available for relocation"],
            Languages=["English"],
            Certifications=[CertificateEntry(Name="Cloud Architect", Issuer="Cloud Org", Year="2024")],
            WorkExperiences=[
                WorkExperience(
                    Company="Analytical Engines",
                    Title="Engineer",
                    Summary="Built systems",
                    RealisationDescription=["Delivered a working engine"],
                )
            ],
            Education=[
                Education(
                    Institution="University",
                    Location="London",
                    Degree="BSc",
                    Field="Mathematics",
                    Year="1843",
                    Completed=True,
                )
            ],
            TotalYearsOfExperience=5,
            PreferredJobTitles=["Staff Engineer"],
            PreferredLocations=["Remote"],
            Skills=[SkillEntry(Name="Python", Level="Advanced")],
        )

    monkeypatch.setattr(profile_builder_router, "extract_profile", complete_profile)
    upload = client.post(
        f"/api/v1/profile/sessions/{session['id']}/files",
        headers=headers,
        files={"file": ("resume.txt", b"complete resume", "text/plain")},
    )

    assert upload.status_code == 202
    profile = client.get("/api/v1/profile", headers=headers)
    data = profile.json()["data"]
    assert data["email"] == "ada@example.com"
    assert data["certifications"] == [{"Name": "Cloud Architect", "Issuer": "Cloud Org", "Year": "2024"}]
    assert data["work_experiences"] == [
        {
            "Company": "Analytical Engines",
            "Title": "Engineer",
            "StartDate": None,
            "EndDate": None,
            "Summary": "Built systems",
            "RealisationDescription": ["Delivered a working engine"],
        }
    ]
    assert data["education"] == [
        {
            "Institution": "University",
            "Location": "London",
            "Degree": "BSc",
            "Field": "Mathematics",
            "Year": "1843",
            "Completed": True,
        }
    ]
    assert data["total_years_of_experience"] == 5
    assert data["preferred_job_titles"] == ["Staff Engineer"]
    assert data["skills"] == [{"Name": "Python", "Level": "Advanced"}]
    assert data["core_competences"] == ["Technical leadership"]
    assert data["additional_informations"] == ["Available for relocation"]


def test_profile_provider_failure_marks_file_and_job_failed(client, monkeypatch):
    token = client.post(
        "/api/v1/auth/register", json={"email": "profile-failure@example.com", "password": "password123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    session = client.post("/api/v1/profile/sessions", headers=headers).json()

    async def fail_profile(*args, **kwargs):
        raise StructuredProviderError("provider unavailable")

    monkeypatch.setattr(profile_builder_router, "extract_profile", fail_profile)
    upload = client.post(
        f"/api/v1/profile/sessions/{session['id']}/files", headers=headers,
        files={"file": ("resume.txt", b"Python engineer", "text/plain")},
    ).json()
    status = client.get(f"/api/v1/processing-jobs/{upload['processing_job_id']}", headers=headers).json()
    detail = client.get(f"/api/v1/profile/sessions/{session['id']}", headers=headers).json()

    assert status["status"] == "failed"
    assert detail["status"] == "failed"
    assert detail["files"][0]["status"] == "failed"


def test_profile_invalid_structured_output_marks_processing_failed(client, monkeypatch):
    token = client.post(
        "/api/v1/auth/register", json={"email": "profile-invalid@example.com", "password": "password123"}
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    session = client.post("/api/v1/profile/sessions", headers=headers).json()

    async def invalid_profile(*args, **kwargs):
        raise StructuredProviderError("invalid structured output")

    monkeypatch.setattr(profile_builder_router, "extract_profile", invalid_profile)
    upload = client.post(
        f"/api/v1/profile/sessions/{session['id']}/files", headers=headers,
        files={"file": ("resume.txt", b"Python engineer", "text/plain")},
    ).json()

    assert client.get(f"/api/v1/processing-jobs/{upload['processing_job_id']}", headers=headers).json()["status"] == "failed"
    assert client.get("/api/v1/profile", headers=headers).status_code == 404
