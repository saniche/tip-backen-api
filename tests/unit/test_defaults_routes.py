import pytest

from shared.database import TestingSessionLocal
from shared.models import UserDocument

PATHS = ["/api/v1/defaults/cv", "/api/v1/defaults/cover-letter"]


def _auth(client, email):
    token = client.post("/api/v1/auth/register", json={"email": email, "password": "password123"}).json()[
        "access_token"
    ]
    return {"Authorization": f"Bearer {token}"}


@pytest.mark.parametrize("path", PATHS)
def test_default_document_add_update_read_delete(client, path):
    headers = _auth(client, f"defaults{path[-5:]}@example.com")

    assert client.get(path, headers=headers).status_code == 404

    created = client.put(path, headers=headers, json={"content": "First version", "filename": "mine.md"})
    assert created.status_code == 200 and created.json()["content"] == "First version"

    updated = client.put(path, headers=headers, json={"content": "Second version"})
    assert updated.status_code == 200
    assert client.get(path, headers=headers).json()["content"] == "Second version"
    assert client.get(path, headers=headers).json()["filename"] is None

    assert client.delete(path, headers=headers).status_code == 204
    assert client.get(path, headers=headers).status_code == 404
    assert client.delete(path, headers=headers).status_code == 404


def test_only_one_document_per_kind_per_user_and_kinds_are_independent(client):
    headers = _auth(client, "one-each@example.com")

    for _ in range(3):
        client.put(PATHS[0], headers=headers, json={"content": "cv"})
    client.put(PATHS[1], headers=headers, json={"content": "letter"})

    with TestingSessionLocal() as db:
        kinds = sorted(item.kind for item in db.query(UserDocument).all())
    assert kinds == ["cover_letter", "cv"]
    assert client.get(PATHS[0], headers=headers).json()["content"] == "cv"
    assert client.get(PATHS[1], headers=headers).json()["content"] == "letter"


def test_default_documents_are_private_to_their_owner(client):
    owner, other = _auth(client, "owner-doc@example.com"), _auth(client, "other-doc@example.com")
    client.put(PATHS[0], headers=owner, json={"content": "secret cv"})

    assert client.get(PATHS[0], headers=other).status_code == 404
    assert client.delete(PATHS[0], headers=other).status_code == 404
    client.put(PATHS[0], headers=other, json={"content": "their own"})
    assert client.get(PATHS[0], headers=owner).json()["content"] == "secret cv"


@pytest.mark.parametrize("body", [{"content": ""}, {"content": "   "}, {"content": "x" * 50_001}, {"nope": 1}])
def test_default_document_rejects_invalid_bodies(client, body):
    headers = _auth(client, "invalid-doc@example.com")

    assert client.put(PATHS[0], headers=headers, json=body).status_code == 422


def test_default_documents_require_authentication(client):
    assert client.get(PATHS[0]).status_code in (401, 403)


# --- file upload (.docx / .md / .txt) ---

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
UPLOAD_PATHS = ["/api/v1/defaults/cv/file", "/api/v1/defaults/cover-letter/file"]


def _docx_bytes(*paragraphs, heading=None):
    from io import BytesIO

    from docx import Document

    document = Document()
    if heading:
        document.add_heading(heading, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize("path", UPLOAD_PATHS)
def test_docx_upload_stores_structured_text_and_replaces_the_previous_document(client, path):
    headers = _auth(client, f"upload{path[-9:-5]}@example.com")
    json_path = path[: -len("/file")]
    client.put(json_path, headers=headers, json={"content": "typed by hand"})

    response = client.post(
        path,
        headers=headers,
        files={"file": ("C:\\docs\\my cv.docx", _docx_bytes("Hello there", heading="Jane"), DOCX)},
    )

    assert response.status_code == 200
    assert response.json()["content"] == "# Jane\nHello there"
    assert response.json()["filename"] == "my cv.docx"
    assert client.get(json_path, headers=headers).json()["content"] == "# Jane\nHello there"
    with TestingSessionLocal() as db:
        assert db.query(UserDocument).count() == 1


def test_markdown_and_text_uploads_work_even_with_a_generic_content_type(client):
    headers = _auth(client, "upload-text@example.com")

    markdown = client.post(
        UPLOAD_PATHS[0],
        headers=headers,
        files={"file": ("cv.md", "\ufeff# Éric\n\nBonjour".encode("utf-8"), "application/octet-stream")},
    )
    plain = client.post(UPLOAD_PATHS[1], headers=headers, files={"file": ("letter.txt", b"Dear team", "text/plain")})

    assert markdown.status_code == 200 and markdown.json()["content"] == "# Éric\n\nBonjour"
    assert plain.status_code == 200 and plain.json()["content"] == "Dear team"


@pytest.mark.parametrize(
    ("filename", "content_type", "body", "message"),
    [
        ("cv.pdf", "application/pdf", b"%PDF-1.4", "Unsupported file type"),
        ("cv.exe", "application/octet-stream", b"MZ", "Unsupported file type"),
        ("cv.docx", DOCX, b"not a zip", "Invalid DOCX file"),
        ("cv.md", "text/markdown", "é".encode("latin-1"), "UTF-8"),
        ("cv.md", "text/markdown", b"   \n ", "No text could be extracted"),
        ("cv.md", "text/markdown", lambda: b"x" * 50_001, "too long"),
        ("cv.md", "text/markdown", lambda: b"x" * (10 * 1024 * 1024 + 1), "maximum size"),
    ],
)
def test_invalid_uploads_are_rejected_and_do_not_replace_the_default(client, filename, content_type, body, message):
    headers = _auth(client, f"bad-upload-{abs(hash((filename, message)))}@example.com")
    client.put(UPLOAD_PATHS[0][: -len("/file")], headers=headers, json={"content": "keep me"})

    content = body() if callable(body) else body
    response = client.post(UPLOAD_PATHS[0], headers=headers, files={"file": (filename, content, content_type)})

    assert response.status_code == 422 and message in response.json()["error"]["message"]
    assert client.get(UPLOAD_PATHS[0][: -len("/file")], headers=headers).json()["content"] == "keep me"


def test_upload_requires_authentication_and_a_file(client):
    assert client.post(UPLOAD_PATHS[0], files={"file": ("a.md", b"x", "text/markdown")}).status_code in (401, 403)
    headers = _auth(client, "no-file@example.com")
    assert client.post(UPLOAD_PATHS[0], headers=headers).status_code == 422
