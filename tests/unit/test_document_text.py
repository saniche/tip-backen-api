from io import BytesIO

import pytest
from docx import Document

from shared.document_text import (
    DOCX_CONTENT_TYPE,
    DocumentTextError,
    detect_document_type,
    extract_document_text,
    extract_docx_text,
    safe_filename,
)


def _docx(build) -> bytes:
    document = Document()
    build(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _sample(document):
    document.add_heading("Jane Doe", level=1)
    document.add_paragraph("Python engineer")
    document.add_heading("Experience", level=2)
    document.add_paragraph("Built APIs", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "FastAPI"
    document.add_paragraph("After the table")


def test_structure_mode_keeps_order_headings_bullets_and_tables():
    text = extract_docx_text(_docx(_sample), preserve_structure=True)

    assert text.splitlines() == [
        "# Jane Doe",
        "Python engineer",
        "",
        "## Experience",
        "- Built APIs",
        "Python | FastAPI",
        "After the table",
    ]


def test_default_mode_keeps_the_profile_builder_behaviour():
    text = extract_docx_text(_docx(_sample))

    assert text.splitlines()[-1] == "Python\tFastAPI"  # tables after paragraphs, tab separated, no markdown
    assert "#" not in text and "- Built APIs" not in text


def test_merged_table_cells_are_not_repeated():
    def build(document):
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"

    assert extract_docx_text(_docx(build), preserve_structure=True) == "Merged"


def test_invalid_docx_is_rejected():
    with pytest.raises(DocumentTextError, match="Invalid DOCX file"):
        extract_docx_text(b"not a zip")


def test_text_files_are_decoded_as_utf8_and_a_bom_is_dropped():
    assert extract_document_text("\ufeffBonjour — Éric".encode("utf-8"), "text/markdown") == "Bonjour — Éric"
    with pytest.raises(DocumentTextError, match="UTF-8"):
        extract_document_text("é".encode("latin-1"), "text/plain")


def test_document_type_detection_prefers_a_known_content_type_then_the_extension():
    assert detect_document_type("text/markdown; charset=utf-8", "x.bin") == "text/markdown"
    assert detect_document_type(DOCX_CONTENT_TYPE, None) == DOCX_CONTENT_TYPE
    assert detect_document_type("application/octet-stream", "cv.DOCX") == DOCX_CONTENT_TYPE
    assert detect_document_type(None, "notes.md") == "text/markdown"
    assert detect_document_type("", "notes.txt") == "text/plain"
    assert detect_document_type("application/octet-stream", "cv.exe") is None
    assert detect_document_type("application/pdf", "cv.docx") is None


def test_safe_filename_drops_client_paths():
    assert safe_filename("C:\\Users\\me\\cv.docx") == "cv.docx"
    assert safe_filename("../../etc/cv.md") == "cv.md"
    assert safe_filename("") is None and safe_filename(None) is None
