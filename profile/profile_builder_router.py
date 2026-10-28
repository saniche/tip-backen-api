from dataclasses import asdict
from io import BytesIO
from profile.profile_builder import extract_profile
from profile.profile_service import apply_profile_edits, merge_profile_values
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, UploadFile
from lxml.etree import XMLSyntaxError
from pydantic import BaseModel
from sqlalchemy import desc, select
from sqlalchemy.orm import Session

from auth.security import get_current_user
from shared.database import SessionLocal, get_db
from shared.models import (
    ProcessingJobType,
    ProfileFileStatus,
    ProfileFragment,
    ProfileSession,
    ProfileSessionStatus,
    User,
    UserFile,
    UserProfile,
)
from shared.processing_service import (
    complete_processing_job,
    create_processing_job,
    fail_processing_job,
    start_processing_job,
)

router = APIRouter(prefix="", tags=["profile"])

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SUPPORTED_PROFILE_CONTENT_TYPES = {
    "application/pdf",
    DOCX_CONTENT_TYPE,
    "text/plain",
    "text/markdown",
    "text/x-markdown",
}


class ProfileUpdate(BaseModel):
    data: dict
    output_language: str = "English"


class ProfileSessionCreate(BaseModel):
    output_language: str = "English"


def _extract_docx_text(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
    except (BadZipFile, KeyError, PackageNotFoundError, ValueError, XMLSyntaxError) as error:
        raise HTTPException(422, "Invalid DOCX file") from error

    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    table_rows = [
        "\t".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join(paragraphs + table_rows)


async def _run_profile_session(session_id: str, user_id: str, processing_job_id: str) -> None:
    db = SessionLocal()
    try:
        session = db.get(ProfileSession, session_id)
        start_processing_job(db, processing_job_id)
        session.status = ProfileSessionStatus.PROCESSING
        db.commit()
        files = db.execute(select(UserFile).where(UserFile.session_id == session_id)).scalars().all()
        fragments = []
        for file in files:
            file.status = ProfileFileStatus.PROCESSING
            extracted = await extract_profile(
                {"content": file.content, "name": file.filename},
                output_language=session.output_language,
            )
            fragment_data = {
                "name": extracted.Name,
                "email": extracted.Email,
                "phone": extracted.Phone,
                "location": extracted.Location,
                "linkedin": extracted.LinkedIn,
                "summary": extracted.Summary,
                "soft_skills": extracted.SoftSkills,
                "core_competences": extracted.CoreCompetences,
                "additional_informations": extracted.AdditionalInformations,
                "certifications": [asdict(item) for item in extracted.Certifications],
                "work_experiences": [asdict(item) for item in extracted.WorkExperiences],
                "education": [asdict(item) for item in extracted.Education],
                "total_years_of_experience": extracted.TotalYearsOfExperience,
                "preferred_job_titles": extracted.PreferredJobTitles,
                "preferred_locations": extracted.PreferredLocations,
                "skills": [asdict(item) for item in extracted.Skills],
                "languages": extracted.Languages,
                "source": file.filename,
            }
            db.add(ProfileFragment(file_id=file.id, data=fragment_data, evidence_type="extracted"))
            fragments.append({"data": fragment_data, "evidence_type": "extracted"})
            file.status = ProfileFileStatus.COMPLETED
        previous = (
            db.execute(select(UserProfile).where(UserProfile.user_id == user_id,
            UserProfile.deleted_at.is_(None),
        ).order_by(desc(UserProfile.created_at)))
            .scalars()
            .first()
        )
        previous_data = previous.data if previous else {}
        previous_evidence = previous.evidence if previous else {}
        merged, evidence = merge_profile_values(previous_data, previous_evidence, fragments)
        profile = UserProfile(user_id=user_id, data=merged, evidence=evidence, output_language=session.output_language)
        db.add(profile)
        db.flush()
        session.status = ProfileSessionStatus.COMPLETED
        db.commit()
        complete_processing_job(db, processing_job_id, profile.id)
    except Exception as ex:
        db.rollback()
        session = db.get(ProfileSession, session_id)
        if session:
            session.status = ProfileSessionStatus.FAILED
        for file in db.execute(select(UserFile).where(UserFile.session_id == session_id)).scalars():
            if file.status != ProfileFileStatus.COMPLETED:
                file.status = ProfileFileStatus.FAILED
                file.error = f"Profile processing failed: {str(ex)}"
        db.commit()
        fail_processing_job(db, processing_job_id, "Profile processing failed")
    finally:
        db.close()


@router.post("/profile/sessions", status_code=201)
def create_profile_session(
    payload: ProfileSessionCreate = ProfileSessionCreate(),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    session = ProfileSession(user_id=current_user.id, output_language=payload.output_language)
    db.add(session)
    db.commit()
    db.refresh(session)
    return {"id": session.id, "status": session.status.value, "output_language": session.output_language}


@router.post("/profile/sessions/{session_id}/files", status_code=202)
async def upload_profile_file(
    session_id: str,
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    session = db.get(ProfileSession, session_id)
    if not session or session.user_id != current_user.id:
        raise HTTPException(404, "Profile session not found")
    content_type = (file.content_type or "").split(";", 1)[0].strip().lower()
    if content_type not in SUPPORTED_PROFILE_CONTENT_TYPES:
        raise HTTPException(422, "Unsupported file type")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(422, "File exceeds maximum size")
    text_content = _extract_docx_text(content) if content_type == DOCX_CONTENT_TYPE else content.decode(
        "utf-8", errors="replace"
    )
    user_file = UserFile(
        user_id=current_user.id,
        session_id=session.id,
        filename=file.filename or "document",
        content_type=file.content_type,
        content=text_content,
    )
    db.add(user_file)
    processing_job = create_processing_job(db, current_user.id, ProcessingJobType.PROFILE_BUILD)
    db.commit()
    background_tasks.add_task(_run_profile_session, session.id, current_user.id, processing_job.id)
    return {"file_id": user_file.id, "processing_job_id": processing_job.id, "status": processing_job.status.value}


@router.get("/profile/sessions/{session_id}")
def get_profile_session(session_id: str, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    session = db.get(ProfileSession, session_id)
    if not session or session.user_id != current_user.id:
        raise HTTPException(404, "Profile session not found")
    files = db.execute(select(UserFile).where(UserFile.session_id == session.id)).scalars().all()
    return {
        "id": session.id,
        "status": session.status.value,
        "output_language": session.output_language,
        "files": [{"id": item.id, "status": item.status.value, "error": item.error} for item in files],
    }


@router.get("/profile")
def get_profile(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    profile = (
        db.execute(
            select(UserProfile).where(UserProfile.user_id == current_user.id,
            UserProfile.deleted_at.is_(None),
        ).order_by(desc(UserProfile.created_at))
        )
        .scalars()
        .first()
    )
    if profile is None:
        raise HTTPException(status_code=404, detail="Profile not found")
    return {
        "id": profile.id,
        "user_id": profile.user_id,
        "data": profile.data,
        "output_language": profile.output_language,
    }


@router.put("/profile")
def update_profile(
    payload: ProfileUpdate, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)
):
    profile = (
        db.execute(
            select(UserProfile).where(UserProfile.user_id == current_user.id,
            UserProfile.deleted_at.is_(None),
        ).order_by(desc(UserProfile.created_at))
        )
        .scalars()
        .first()
    )
    if profile is None:
        profile = UserProfile(user_id=current_user.id, data={}, evidence={}, output_language=payload.output_language)
        db.add(profile)
    apply_profile_edits(profile, payload.data, payload.output_language, None)
    db.commit()
    db.refresh(profile)
    return {
        "id": profile.id,
        "user_id": profile.user_id,
        "data": profile.data,
        "output_language": profile.output_language,
    }
