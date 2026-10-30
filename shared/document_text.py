"""Plain-text extraction for uploaded documents (.docx, .md, .txt)."""

import os
import re
from io import BytesIO
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml.etree import XMLSyntaxError

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_CONTENT_TYPES = {"text/plain", "text/markdown", "text/x-markdown"}
_GENERIC_CONTENT_TYPES = {"", "application/octet-stream"}
_EXTENSION_TYPES = {
    ".docx": DOCX_CONTENT_TYPE,
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".txt": "text/plain",
}


class DocumentTextError(ValueError):
    """The uploaded document cannot be turned into text."""


def detect_document_type(content_type: str | None, filename: str | None) -> str | None:
    """Supported content type for an upload, or None.

    A recognised content type wins. Browsers and HTTP clients often label .docx/.md files as
    application/octet-stream (or nothing), so only then the file extension decides.
    """
    declared = (content_type or "").split(";", 1)[0].strip().lower()
    if declared == DOCX_CONTENT_TYPE or declared in TEXT_CONTENT_TYPES:
        return declared
    if declared in _GENERIC_CONTENT_TYPES:
        return _EXTENSION_TYPES.get(os.path.splitext(filename or "")[1].lower())
    return None


def safe_filename(filename: str | None, max_length: int = 200) -> str | None:
    """Base name only (no client-side path), or None."""
    name = os.path.basename((filename or "").replace("\\", "/")).strip()
    return name[:max_length] or None


def _iter_blocks(document):
    """Paragraphs and tables in document order."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _paragraph_line(paragraph) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    style = (getattr(paragraph.style, "name", "") or "").strip()
    if style == "Title":
        return f"# {text}"
    heading = re.fullmatch(r"Heading (\d)", style)
    if heading:
        return f"{'#' * min(max(int(heading.group(1)), 1), 6)} {text}"
    if style.startswith("List"):
        return f"- {text}"
    return text


def _table_lines(table) -> list[str]:
    lines = []
    for row in table.rows:
        seen, cells = set(), []
        for cell in row.cells:
            if id(cell._tc) in seen:  # merged cells repeat the same underlying cell
                continue
            seen.add(id(cell._tc))
            text = " ".join(cell.text.split())
            if text:
                cells.append(text)
        if cells:
            lines.append(" | ".join(cells))
    return lines


def extract_docx_text(content: bytes, preserve_structure: bool = False) -> str:
    """Text of a .docx.

    Default (used by the profile builder): paragraphs, then table rows. With preserve_structure the
    text follows the document order and keeps headings and bullet points as markdown, which is
    better for a style reference.
    """
    try:
        document = Document(BytesIO(content))
    except (BadZipFile, KeyError, PackageNotFoundError, ValueError, XMLSyntaxError) as error:
        raise DocumentTextError("Invalid DOCX file") from error

    if not preserve_structure:
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        table_rows = ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
        return "\n".join(paragraphs + table_rows)

    lines: list[str] = []
    for block in _iter_blocks(document):
        block_lines = _table_lines(block) if isinstance(block, Table) else [_paragraph_line(block)]
        for line in block_lines:
            if not line:
                continue
            if line.startswith("#") and lines and lines[-1] != "":
                lines.append("")  # blank line before a heading
            lines.append(line)
    return "\n".join(lines)


def extract_document_text(content: bytes, content_type: str, preserve_structure: bool = False) -> str:
    if content_type == DOCX_CONTENT_TYPE:
        return extract_docx_text(content, preserve_structure)
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise DocumentTextError("The file must be UTF-8 encoded text") from error
